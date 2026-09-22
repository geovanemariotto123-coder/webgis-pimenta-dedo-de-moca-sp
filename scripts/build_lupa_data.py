#!/usr/bin/env python3
"""Baixa e consolida a categoria Pimenta dos arquivos municipais do LUPA."""

import json
import gzip
import re
import unicodedata
from concurrent.futures import ThreadPoolExecutor, as_completed
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from urllib.parse import urljoin
from urllib.request import Request, urlopen

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "lupa_pimenta.json"
HEADERS = {"User-Agent": "GeoTerra-WebGIS/1.0"}
IBGE_URL = "https://servicodados.ibge.gov.br/api/v1/localidades/estados/35/municipios"
EDITIONS = {
    "2007/08": "https://www.cati.sp.gov.br/projetolupa/dadosmunicipais.php",
    "2016/17": "https://www.cati.sp.gov.br/projetolupa/dadosmunicipais1617.php",
}


def fetch(url):
    request = Request(url, headers=HEADERS)
    with urlopen(request, timeout=90) as response:
        data = response.read()
        return gzip.decompress(data) if data[:2] == b"\x1f\x8b" else data


def normalize(value):
    value = unicodedata.normalize("NFKD", str(value or ""))
    value = "".join(char for char in value if not unicodedata.combining(char))
    return re.sub(r"[^A-Z0-9]+", " ", value.upper()).strip()


class LinkParser(HTMLParser):
    def __init__(self, base_url):
        super().__init__()
        self.base_url = base_url
        self.href = None
        self.links = []

    def handle_starttag(self, tag, attrs):
        if tag == "a":
            self.href = dict(attrs).get("href")

    def handle_data(self, data):
        name = data.strip()
        if self.href and name and self.href.lower().endswith(".xlsx"):
            self.links.append((name, urljoin(self.base_url, self.href)))

    def handle_endtag(self, tag):
        if tag == "a":
            self.href = None


def edition_links(url):
    parser = LinkParser(url)
    parser.feed(fetch(url).decode("utf-8", "replace"))
    return parser.links


def parse_workbook(name, url):
    workbook = load_workbook(BytesIO(fetch(url)), data_only=True, read_only=True)
    worksheet = workbook.active
    for row in worksheet.iter_rows(values_only=True):
        if normalize(row[0]) == "PIMENTA":
            return {
                "name": name,
                "upas": int(row[1]),
                "area_ha": round(float(row[5]), 2),
            }
    return None


def main():
    municipalities = json.loads(fetch(IBGE_URL))
    codes = {normalize(item["nome"]): str(item["id"]) for item in municipalities}
    output = {
        "title": "Pimenta — todas as variedades",
        "source": "CATI/IEA, Projeto LUPA",
        "warning": "A categoria Pimenta não identifica variedades e não representa exclusivamente dedo-de-moça.",
        "metrics": {
            "area_ha": ["Área cultivada", "ha"],
            "upas": ["Unidades de produção agropecuária", "UPAs"],
        },
        "years": {},
    }
    for edition, page_url in EDITIONS.items():
        links = edition_links(page_url)
        rows = []
        with ThreadPoolExecutor(max_workers=16) as executor:
            futures = {executor.submit(parse_workbook, name, url): (name, url) for name, url in links}
            for future in as_completed(futures):
                row = future.result()
                if not row:
                    continue
                code = codes.get(normalize(row["name"]))
                if not code:
                    raise RuntimeError(f"Município sem correspondência no IBGE: {row['name']}")
                row["code"] = code
                rows.append(row)
        rows.sort(key=lambda item: item["name"])
        output["years"][edition] = {
            "source_url": page_url,
            "municipalities_with_records": len(rows),
            "total_area_ha": round(sum(row["area_ha"] for row in rows), 2),
            "total_upas": sum(row["upas"] for row in rows),
            "values": {row.pop("code"): row for row in rows},
        }
        print(
            edition,
            f"{len(rows)} municípios",
            f"{output['years'][edition]['total_area_ha']:.2f} ha",
            f"{output['years'][edition]['total_upas']} UPAs",
        )
    OUTPUT.write_text(json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
